# -*- coding: utf-8 -*-
import hashlib, json, os, sys
from pptx import Presentation
from pptx.util import Emu

SRC = r"D:\Java\Obisidian笔记\AI"
OUT = r"D:\Java\Obisidian笔记\.claude\plans\2026-10-25-jvm-juc-ai\ai_work"
FILES = {
    "01-提示词工程": "01提示词工程.pptx",
    "02-大模型RAG开发": "02大模型RAG开发.pptx",
    "03-RAG项目": "03RAG项目.pptx",
    "04-LangChain4J与通义千问": "LangChain4J+通义千问大模型应用开发(志愿填报大师).pptx",
}

def walk(shapes, out, slide, prefix=""):
    for sh in shapes:
        try:
            st = sh.shape_type
        except Exception:
            st = None
        name = sh.name
        info = {"name": name, "type": str(st)}
        if sh.shape_type == 6:  # group
            try:
                walk(sh.shapes, out, slide, prefix + name + "/")
            except Exception:
                pass
        if sh.has_text_frame:
            txt = "\n".join(p.text for p in sh.text_frame.paragraphs)
            if txt.strip():
                info["text"] = txt
                info["level"] = [p.level for p in sh.text_frame.paragraphs]
                info["top"] = round(Emu(sh.top).inches, 2) if sh.top is not None else None
                info["left"] = round(Emu(sh.left).inches, 2) if sh.left is not None else None
                out["texts"].append(info)
        if getattr(sh, "has_table", False) and sh.has_table:
            rows = []
            for r in sh.table.rows:
                rows.append([c.text for c in r.cells])
            out["tables"].append({"name": name, "rows": rows})
        if sh.shape_type == 13:  # picture
            try:
                img = sh.image
                blob = img.blob
                sha = hashlib.sha1(blob).hexdigest()
                ext = img.ext
                fn = f"{slide:03d}-{len(out['images'])+1:02d}.{ext}"
                rec = {
                    "idx": len(out["images"]) + 1,
                    "filename": fn,
                    "sha1": sha,
                    "ext": ext,
                    "bytes": len(blob),
                    "px_w": img.size[0], "px_h": img.size[1],
                    "shape_w_in": round(Emu(sh.width).inches, 2),
                    "shape_h_in": round(Emu(sh.height).inches, 2),
                    "top_in": round(Emu(sh.top).inches, 2),
                    "left_in": round(Emu(sh.left).inches, 2),
                    "blob": blob,
                }
                out["images"].append(rec)
            except Exception as e:
                out["errors"].append(f"img {name}: {e}")

allmeta = {}
for note, fn in FILES.items():
    path = os.path.join(SRC, fn)
    prs = Presentation(path)
    slides = []
    W, H = Emu(prs.slide_width).inches, Emu(prs.slide_height).inches
    for i, slide in enumerate(prs.slides, 1):
        out = {"slide": i, "texts": [], "tables": [], "images": [], "notes": "", "errors": []}
        try:
            walk(slide.shapes, out, i)
        except Exception as e:
            out["errors"].append(str(e))
        if slide.has_notes_slide:
            nt = slide.notes_slide.notes_text_frame.text
            out["notes"] = nt.strip()
        slides.append(out)
    allmeta[note] = {"file": fn, "count": len(slides), "w": W, "h": H, "slides": slides}
    # save blobs
    for s in slides:
        for im in s["images"]:
            d = os.path.join(OUT, "images", note)
            os.makedirs(d, exist_ok=True)
            with open(os.path.join(d, im["filename"]), "wb") as f:
                f.write(im["blob"])
            del im["blob"]
    with open(os.path.join(OUT, "json", note + ".json"), "w", encoding="utf-8") as f:
        json.dump(allmeta[note], f, ensure_ascii=False, indent=1)
    # text dump
    lines = [f"# {note} ({fn}) slides={len(slides)} size={W}x{H}"]
    for s in slides:
        lines.append(f"\n===== SLIDE {s['slide']} =====")
        for t in s["texts"]:
            lv = ",".join(map(str, t.get("level", [])))
            lines.append(f"[{t['name']}] lv={lv} pos=({t.get('left')},{t.get('top')})")
            for ln in t["text"].split("\n"):
                lines.append("  " + ln)
        for tb in s["tables"]:
            lines.append(f"[TABLE {tb['name']}]")
            for r in tb["rows"]:
                lines.append(" | ".join(r))
        for im in s["images"]:
            lines.append(f"[IMG {im['idx']:02d}] {im['filename']} {im['px_w']}x{im['px_h']}px {im['bytes']}B shape={im['shape_w_in']}x{im['shape_h_in']}in pos=({im['left_in']},{im['top_in']}) sha={im['sha1'][:10]}")
        if s["notes"]:
            lines.append("[NOTES] " + s["notes"].replace("\n", "\n  "))
        if s["errors"]:
            lines.append("[ERROR] " + "; ".join(s["errors"]))
    with open(os.path.join(OUT, "text", note + ".txt"), "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    print(note, "slides", len(slides), "images", sum(len(s["images"]) for s in slides), "tables", sum(len(s["tables"]) for s in slides), "notes", sum(1 for s in slides if s["notes"]))